import csv
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from trade_buyer.deduplicate import deduplicate
from trade_buyer.exporters import export_excel
from trade_buyer.importers import REQUIRED, import_buyers
from trade_buyer.normalize import classify_materials, normalize_company

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "sample_data/buyers.csv"


class SystemTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def write_csv(self, rows, headers=REQUIRED):
        path = self.directory / "input.csv"
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(rows)
        return path

    def real_row(self):
        # Test-only fixture, never a claimed company or real buyer.
        return ["Fixture Metals Ltd", "Fixture Country", "copper scrap", "Fixture source",
                "https://example.invalid/source", "2026-01-01", "permission_granted", "real"]

    def test_normalization(self):
        self.assertEqual(normalize_company("  ACME & Sons, Limited "), "acme and sons")
        self.assertEqual(normalize_company("ÉCO Metals LLC"), "éco metals")
        self.assertEqual(normalize_company("株式会社 金属"), "株式会社 金属")

    def test_material_classification(self):
        self.assertEqual(classify_materials("aluminium scrap; copper scrap; aluminum ingots"),
                         {"aluminum_scrap", "copper_scrap", "metal_ingot"})
        self.assertEqual(classify_materials("copper wire"), {"unclassified"})
        self.assertEqual(classify_materials("notingots"), {"unclassified"})

    def test_sample_opt_in(self):
        buyers, rejected = import_buyers(SAMPLE)
        self.assertEqual((len(buyers), rejected), (0, 3))

    def test_merge_preserves_sources_and_input(self):
        records, rejected = import_buyers(SAMPLE, True)
        buyers, duplicates = deduplicate(records)
        self.assertEqual((len(buyers), duplicates, rejected), (2, 1, 0))
        self.assertEqual(len(buyers[0].sources), 2)
        self.assertEqual(buyers[0].materials, {"aluminum_scrap", "copper_scrap"})
        self.assertEqual(len(records[0].sources), 1)

    def test_country_prevents_merge(self):
        row = self.real_row()
        other = row.copy()
        other[1] = "Other Country"
        buyers, _ = import_buyers(self.write_csv([row, other]))
        self.assertEqual(deduplicate(buyers)[1], 0)

    def test_rejections_logged(self):
        rows = []
        for index, value in [(5, "bad-date"), (6, "unknown"), (4, "file:///local"), (0, "!!!"), (1, "")]:
            row = self.real_row()
            row[index] = value
            rows.append(row)
        with self.assertLogs(level="WARNING") as captured:
            buyers, rejected = import_buyers(self.write_csv(rows))
        self.assertEqual((len(buyers), rejected), (0, 5))
        self.assertEqual(len(captured.output), 5)

    def test_invalid_schema(self):
        with self.assertRaises(ValueError):
            import_buyers(self.write_csv([], ["company_name"]))
        with self.assertRaises(ValueError):
            import_buyers(self.write_csv([], [*REQUIRED, "country"]))

    def test_excel_import(self):
        path = self.directory / "input.xlsx"
        book = Workbook()
        book.active.append(list(REQUIRED))
        book.active.append(self.real_row())
        book.save(path)
        book.close()
        buyers, rejected = import_buyers(path)
        self.assertEqual((len(buyers), rejected), (1, 0))
        self.assertEqual(buyers[0].sources[0].authorization, "permission_granted")

    def test_export_provenance_and_formula_protection(self):
        row = self.real_row()
        row[0] = "=Fixture Ltd"
        row[3] = "@Fixture source"
        buyers, _ = import_buyers(self.write_csv([row]))
        path = self.directory / "report.xlsx"
        export_excel(buyers, path)
        book = load_workbook(path)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["Buyers", "Sources"])
        self.assertEqual(book["Buyers"].max_row, 2)
        self.assertEqual(book["Sources"]["B2"].value, "'@Fixture source")
        self.assertEqual(book["Buyers"]["B2"].data_type, "s")
        self.assertEqual(book["Sources"]["E2"].value, "permission_granted")

    def run_cli(self, *arguments):
        return subprocess.run([sys.executable, "-m", "trade_buyer", *map(str, arguments)],
                              cwd=ROOT, capture_output=True, text=True)

    def test_cli_end_to_end(self):
        output = self.directory / "report.xlsx"
        result = self.run_cli(SAMPLE, "--allow-samples", "--output", output,
                              "--log-file", self.directory / "errors.log")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Exported 2 buyers; merged 1 duplicates", result.stdout)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book["Sources"].max_row, 4)

    def test_cli_partial_success(self):
        row = self.real_row()
        invalid = row.copy()
        invalid[5] = "invalid"
        path = self.write_csv([row, invalid])
        result = self.run_cli(path, "--output", self.directory / "partial.xlsx",
                              "--log-file", self.directory / "errors.log")
        self.assertEqual(result.returncode, 2)
        self.assertIn("rejected 1 rows", result.stdout)
        self.assertIn("row 3 rejected", (self.directory / "errors.log").read_text())

    def test_cli_missing_input_and_overwrite(self):
        result = self.run_cli(self.directory / "missing.csv", "--output", self.directory / "out.xlsx",
                              "--log-file", self.directory / "errors.log")
        self.assertEqual(result.returncode, 1)
        path = self.directory / "input.xlsx"
        path.write_bytes(b"preserve me")
        result = self.run_cli(path, "--output", path)
        self.assertEqual(result.returncode, 1)
        self.assertEqual(path.read_bytes(), b"preserve me")


if __name__ == "__main__":
    unittest.main()

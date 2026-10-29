"""Offline API contract fixtures: none represent real buyers."""
from datetime import date
import json
import csv
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from io import BytesIO
from email.message import Message
from urllib.error import HTTPError

from openpyxl import load_workbook
from trade_buyer.discovery import discover, fetch_json, main, NoRedirect, MAX_BYTES
from trade_buyer.deduplicate import deduplicate
from trade_buyer.exporters import export_excel


class DiscoveryTests(unittest.TestCase):
    def setUp(self):
        self.config = {"name": "TEST FIXTURE ONLY", "endpoint": "https://example.invalid/api",
                       "terms_url": "https://example.invalid/terms", "permission_basis": "TEST ONLY",
                       "kind": "trade_records", "authorization": "public_license",
                       "automated_access_permitted": True}
        self.record = {"company_name": "TEST FIXTURE Metals Ltd", "country": "TEST COUNTRY",
                       "products_purchased": "aluminum scrap; copper scrap; metal ingots",
                       "source_url": "https://example.invalid/record/1", "buyer_evidence": "TEST importer entry",
                       "record_date": "2026-01-01", "buyer_role": "importer",
                       "business_email": "fixture@example.invalid", "email_public": True,
                       "website": "https://example.invalid"}

    def discover_rows(self, rows):
        return discover(self.config, lambda url: {"records": rows})

    def test_evidence_and_contact_provenance_export(self):
        buyers, rejected = self.discover_rows([self.record])
        self.assertEqual(rejected, 0)
        self.assertEqual(buyers[0].materials, {"aluminum_scrap", "copper_scrap", "metal_ingot"})
        self.assertEqual(buyers[0].sources[0].verified_on, date.today().isoformat())
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "buyers.xlsx"
            export_excel(buyers, output)
            book = load_workbook(output)
            try:
                columns = {cell.value: cell.column for cell in book["Sources"][1]}
                self.assertEqual(book["Sources"].cell(2, columns["business_email"]).value,
                                 "fixture@example.invalid")
                self.assertEqual(book["Sources"].cell(2, columns["buyer_evidence"]).value,
                                 "TEST importer entry")
            finally:
                book.close()

    def test_unconfirmed_lead_not_claimed_as_purchaser(self):
        row = dict(self.record, buyer_role="lead", buyer_evidence="")
        buyers, rejected = self.discover_rows([row])
        self.assertEqual(rejected, 0)
        self.assertEqual(buyers[0].sources[0].verification_status, "unconfirmed_lead")
        self.assertEqual(buyers[0].sources[0].products_purchased, "")

    def test_missing_evidence_downgrades_buyer_to_lead(self):
        buyers, rejected = self.discover_rows([dict(self.record, buyer_evidence="")])
        self.assertEqual(rejected, 0)
        self.assertEqual(buyers[0].sources[0].verification_status, "unconfirmed_lead")

    def test_cli_local_export_without_network(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = root / "source.json"
            records = root / "records.json"
            config.write_text(json.dumps(self.config))
            records.write_text(json.dumps({"records": [self.record]}))
            with patch("trade_buyer.discovery.fetch_json", side_effect=AssertionError("No network")):
                status = main(["--source-config", str(config), "--records-file", str(records),
                               "--output", str(root / "out.xlsx"), "--log-file", str(root / "run.log")])
            self.assertEqual(status, 0)

    def test_csv_contact_import_preserves_provenance_without_trusting_status(self):
        from trade_buyer.importers import import_buyers, REQUIRED
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "contacts.csv"
            columns = [*REQUIRED, "business_email", "email_public", "website", "verification_status"]
            row = ["TEST FIXTURE Ltd", "TEST COUNTRY", "copper scrap", "TEST SOURCE",
                   "https://example.invalid/record", "2026-01-01", "public_license", "real",
                   "contact@example.invalid", "true", "https://example.invalid", "documented_buyer_activity"]
            with path.open("w", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(columns)
                writer.writerow(row)
            buyers, rejected = import_buyers(path)
            self.assertEqual(rejected, 0)
            self.assertEqual(buyers[0].sources[0].business_email, "contact@example.invalid")
            self.assertEqual(buyers[0].sources[0].verification_status, "unconfirmed_lead")
            row[9] = "false"
            with path.open("w", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(columns)
                writer.writerow(row)
            self.assertEqual(import_buyers(path)[1], 1)

    def test_html_response_rejected(self):
        class Response(BytesIO):
            headers = Message()
        Response.headers["Content-Type"] = "text/html"
        with patch("trade_buyer.discovery.urllib.request.build_opener") as opener:
            opener.return_value.open.return_value = Response(b"<html>login</html>")
            with self.assertRaises(ValueError):
                fetch_json(self.config["endpoint"])

    def test_permissions_fail_before_network(self):
        for field, value in [("automated_access_permitted", False), ("authorization", "unknown"),
                             ("terms_url", ""), ("endpoint", "http://example.invalid"),
                             ("endpoint", "https://user:secret@example.invalid"), ("kind", "scrape")]:
            with self.subTest(field=field):
                config = dict(self.config, **{field: value})
                with self.assertRaises(ValueError):
                    discover(config, lambda _: self.fail("Network must not be reached"))

    def test_seller_and_missing_evidence_rejected(self):
        for field, value in [("buyer_role", "seller"),
                             ("products_purchased", "copper wire"), ("record_date", "2999-01-01"),
                             ("email_public", False), ("business_email", "not-an-email"),
                             ("website", "javascript:alert(1)"), ("country", "")]:
            with self.subTest(field=field):
                rows, rejected = self.discover_rows([dict(self.record, **{field: value})])
                self.assertEqual((rows, rejected), ([], 1))

    def test_absent_contacts_not_invented(self):
        row = dict(self.record)
        del row["business_email"]
        del row["website"]
        buyers, _ = self.discover_rows([row])
        self.assertEqual(buyers[0].sources[0].business_email, "")
        self.assertEqual(buyers[0].sources[0].website, "")

    def test_duplicate_sources_preserved(self):
        other = dict(self.record, company_name="TEST FIXTURE Metals Limited",
                     source_url="https://example.invalid/record/2", business_email="second@example.invalid")
        buyers, _ = self.discover_rows([self.record, other])
        merged, count = deduplicate(buyers)
        self.assertEqual(count, 1)
        self.assertEqual(len(merged[0].sources), 2)
        self.assertEqual(merged[0].sources[1].business_email, "second@example.invalid")

    def test_bad_schema_and_pagination(self):
        for payload in [[], {}, {"records": {}}, {"records": [], "next": "https://example.invalid/page2"}]:
            with self.assertRaises(ValueError):
                discover(self.config, lambda _: payload)

    def test_access_denial_propagates_without_retry(self):
        calls = []
        def denied(url):
            calls.append(url)
            raise HTTPError(url, 403, "Forbidden", None, None)
        with self.assertRaises(HTTPError):
            discover(self.config, denied)
        self.assertEqual(len(calls), 1)

    def test_redirect_refused(self):
        with self.assertRaises(ValueError):
            NoRedirect().redirect_request(None, None, 302, "Redirect", {}, "https://other.invalid")

    def test_transport_json_and_bounds(self):
        class Response(BytesIO):
            headers = Message()
        Response.headers["Content-Type"] = "application/json"
        for body, valid in [(b'{"records": []}', True), (b'x' * (MAX_BYTES + 1), False), (b'bad json', False)]:
            with patch("trade_buyer.discovery.urllib.request.build_opener") as opener:
                opener.return_value.open.return_value = Response(body)
                if valid:
                    self.assertEqual(fetch_json(self.config["endpoint"]), {"records": []})
                else:
                    with self.assertRaises(ValueError):
                        fetch_json(self.config["endpoint"])
                self.assertEqual(opener.return_value.open.call_count, 1)

    def test_cli_export_with_mock_api(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = root / "source.json"
            config.write_text(json.dumps(self.config))
            records, _ = self.discover_rows([self.record])
            with patch("trade_buyer.discovery.discover", return_value=(records, 0)):
                status = main(["--source-config", str(config), "--output", str(root / "out.xlsx"),
                               "--log-file", str(root / "run.log")])
            self.assertEqual(status, 0)
            self.assertTrue((root / "out.xlsx").exists())

    def test_cli_zero_records_is_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = root / "source.json"
            config.write_text(json.dumps(self.config))
            with patch("trade_buyer.discovery.discover", return_value=([], 0)):
                status = main(["--source-config", str(config), "--output", str(root / "out.xlsx"),
                               "--log-file", str(root / "run.log")])
            self.assertEqual(status, 1)
            self.assertFalse((root / "out.xlsx").exists())

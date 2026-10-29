"""Local authorized-file ingestion; no network collection is performed."""
import csv
import logging
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook

from .models import Buyer, Source
from .normalize import classify_materials, normalize_company

REQUIRED = ("company_name", "country", "materials", "source_name", "source_url",
            "collected_on", "authorization", "record_type")
AUTHORIZATIONS = {"public_license", "permission_granted", "synthetic_sample"}


def read_rows(path: Path):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            yield from reader
    elif path.suffix.lower() == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            yield from book.active.iter_rows(values_only=True)
        finally:
            book.close()
    else:
        raise ValueError(f"Unsupported input format: {path.suffix}; use .csv or .xlsx")


def import_buyers(path: Path, allow_samples: bool = False) -> tuple[list[Buyer], int]:
    rows = iter(read_rows(path))
    header = [str(cell or "").strip() for cell in next(rows, ())]
    if len(header) != len(set(header)) or not set(REQUIRED).issubset(header):
        raise ValueError("Input requires unique headers: " + ", ".join(REQUIRED))
    buyers, rejected = [], 0
    for number, row in enumerate(rows, 2):
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        try:
            if len(row) != len(header):
                raise ValueError("column count differs from header")
            values = {key: str(value).strip() if value is not None else ""
                      for key, value in zip(header, row)}
            if any(not values[key] for key in REQUIRED):
                raise ValueError("required field is empty")
            normalized = normalize_company(values["company_name"])
            if not normalized:
                raise ValueError("company name has no searchable characters")
            date.fromisoformat(values["collected_on"])
            authorization = values["authorization"]
            kind = values["record_type"]
            if authorization not in AUTHORIZATIONS or kind not in {"real", "synthetic"}:
                raise ValueError("invalid authorization or record_type")
            if kind == "synthetic":
                if not allow_samples or authorization != "synthetic_sample":
                    raise ValueError("synthetic rows require --allow-samples and synthetic_sample authorization")
                if not values["company_name"].startswith("SYNTHETIC SAMPLE"):
                    raise ValueError("synthetic company name must start with SYNTHETIC SAMPLE")
            else:
                parsed = urlparse(values["source_url"])
                if authorization == "synthetic_sample" or parsed.scheme not in {"https", "http"} or not parsed.hostname:
                    raise ValueError("real records require an HTTP(S) source and explicit authorization")
            source = Source(values["source_name"], values["source_url"], values["collected_on"], authorization, kind)
            buyers.append(Buyer(values["company_name"], normalized, values["country"],
                                classify_materials(values["materials"]), [source]))
        except ValueError as error:
            rejected += 1
            logging.warning("%s row %d rejected: %s", path.name, number, error)
    return buyers, rejected
